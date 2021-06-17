from ...contents import Table
from pptx.shapes.shapetree import TablePlaceholder
from ..helpers import df_to_table


create_table = TablePlaceholder.insert_table

def replace_with_table(pptable, shape, slide):
    table = slide.shapes.add_table(pptable.height, pptable.width, shape.left, shape.top, shape.width, shape.height)

    #calculate max width/height for target size
    # ratio = min(shape.width / float(table.width), shape.height / float(pic.height))
    #
    # table.height = shape.height
    # table.width = shape.width
    #
    # table.left = int(shape.left + ((shape.width - table.width) / 2))
    # table.top = int(shape.top + ((shape.height - table.height) / 2))

    placeholder = shape.element
    placeholder.getparent().remove(placeholder)
    return table


class PptTable(Table):

    @classmethod
    def from_parent(cls, parent):
        return cls(parent.styler)

    # def set_params(self, na_rep=None, float_format=None, columns=None, header=None, index=True, index_label=None,
    #                merge_cells=True, inf_rep="inf", show_title=True):
    #     for key in self.params:
    #         if eval(key) is not None:
    #             self.params[key] = eval(key)
    #     return self

    def save(self, placeholder, slide_obj):
        table = replace_with_table(self, placeholder, slide_obj)
        df_to_table(table, self.data)

